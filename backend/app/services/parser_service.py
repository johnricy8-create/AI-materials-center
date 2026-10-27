from __future__ import annotations

import html
import re
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

import pymupdf
from charset_normalizer import from_bytes
from docx import Document
from markdown_it import MarkdownIt
from PIL import Image

MAX_TEXT_CHARS = 200_000


@dataclass
class ParsedContent:
    text: str = ""
    page_count: int | None = None
    width: int | None = None
    height: int | None = None
    thumbnail_bytes: bytes | None = None
    meta: dict = field(default_factory=dict)

    @property
    def char_count(self) -> int:
        return len(self.text)


def _clean_text(value: str) -> str:
    value = value.replace("\x00", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()[:MAX_TEXT_CHARS]


def _image_thumbnail(image: Image.Image) -> bytes:
    thumb = image.convert("RGB").copy()
    thumb.thumbnail((960, 960))
    buffer = BytesIO()
    thumb.save(buffer, format="JPEG", quality=84, optimize=True)
    return buffer.getvalue()


def prepare_image_for_vision(path: Path, max_size: int = 1200) -> bytes:
    with Image.open(path) as image:
        prepared = image.convert("RGB").copy()
        prepared.thumbnail((max_size, max_size))
        buffer = BytesIO()
        prepared.save(buffer, format="JPEG", quality=85, optimize=True)
        return buffer.getvalue()


def _pdf_thumbnail(document: pymupdf.Document) -> bytes | None:
    if document.page_count == 0:
        return None
    page = document.load_page(0)
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.3, 1.3), alpha=False)
    image = Image.open(BytesIO(pixmap.tobytes("png")))
    return _image_thumbnail(image)


def _parse_image(path: Path) -> ParsedContent:
    with Image.open(path) as image:
        width, height = image.size
        return ParsedContent(
            width=width,
            height=height,
            thumbnail_bytes=_image_thumbnail(image),
            meta={"format": image.format or "", "mode": image.mode},
        )


def _parse_pdf(path: Path) -> ParsedContent:
    with pymupdf.open(str(path)) as document:
        if document.needs_pass:
            raise ValueError("PDF 已加密，暂时无法解析")
        text_parts: list[str] = []
        for index in range(min(document.page_count, 200)):
            text_parts.append(document.load_page(index).get_text("text"))
        return ParsedContent(
            text=_clean_text("\n".join(text_parts)),
            page_count=document.page_count,
            thumbnail_bytes=_pdf_thumbnail(document),
            meta={"encrypted": False},
        )


def _parse_docx(path: Path) -> ParsedContent:
    document = Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return ParsedContent(
        text=_clean_text("\n".join(parts)),
        meta={"paragraphs": len(document.paragraphs), "tables": len(document.tables)},
    )


def _decode_text(data: bytes) -> str:
    match = from_bytes(data).best()
    if match is not None:
        return str(match)
    return data.decode("utf-8", errors="replace")


def _parse_text(path: Path, *, markdown: bool) -> ParsedContent:
    raw = _decode_text(path.read_bytes())
    if markdown:
        rendered = MarkdownIt().render(raw)
        raw = html.unescape(re.sub(r"<[^>]+>", " ", rendered))
    return ParsedContent(text=_clean_text(raw), meta={"markdown": markdown})


def parse_file(path: Path, file_type: str) -> ParsedContent:
    if not path.exists():
        raise FileNotFoundError("原文件不存在")

    if file_type == "image":
        return _parse_image(path)
    if file_type == "pdf":
        return _parse_pdf(path)
    if file_type == "docx":
        return _parse_docx(path)
    if file_type == "md":
        return _parse_text(path, markdown=True)
    if file_type == "txt":
        return _parse_text(path, markdown=False)
    raise ValueError(f"暂不支持解析：{file_type}")
